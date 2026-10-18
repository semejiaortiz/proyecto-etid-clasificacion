#!/usr/bin/env python3
"""Convierte el plan ETID en un Word compacto, editable y autocontenido."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

NAVY = "123A6D"
BLUE = "1E5A96"
PALE = "EAF1F8"
LIGHT = "F6F8FB"
TEXT = "202735"
GRAY = "5D6675"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=70, start=85, bottom=70, end=85) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    relation_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE_PATTERN = re.compile(r"(https?://\S+|\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)")


def add_inline(paragraph, text: str) -> None:
    position = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position:match.start()])
        token = match.group(0)
        if token.startswith("http"):
            clean = token.rstrip(".,;)")
            suffix = token[len(clean):]
            add_hyperlink(paragraph, clean, clean)
            if suffix:
                paragraph.add_run(suffix)
        elif token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(8.8)
            run.font.color.rgb = RGBColor.from_string(NAVY)
        else:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def set_repeat_page_number(header) -> None:
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run._r.extend((field_begin, instruction, field_end))
    run.font.name = "Arial"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(GRAY)


def style_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.55)
    section.bottom_margin = Cm(1.55)
    section.left_margin = Cm(1.75)
    section.right_margin = Cm(1.75)
    section.header_distance = Cm(0.7)
    section.footer_distance = Cm(0.7)
    set_repeat_page_number(section.header)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.7)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    normal.paragraph_format.space_after = Pt(3.2)
    normal.paragraph_format.widow_control = True
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")

    for name, size, color, before, after in (
        ("Heading 1", 13.5, NAVY, 8, 3),
        ("Heading 2", 11.2, BLUE, 6, 2.5),
        ("Heading 3", 10.2, BLUE, 4, 2),
    ):
        style = document.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Técnicas Avanzadas de Modelado en IA · Unidad 4")
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(GRAY)


def add_title_block(document: Document, lines: list[str]) -> int:
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(8)
    title.paragraph_format.space_after = Pt(3)
    run = title.add_run(lines[0].removeprefix("# "))
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(19)
    run.font.color.rgb = RGBColor.from_string(NAVY)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(9)
    run = subtitle.add_run(lines[2].removeprefix("## "))
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor.from_string(BLUE)

    index = 4
    metadata = []
    while index < len(lines) and lines[index].strip() != "---":
        if lines[index].strip():
            metadata.append(lines[index].strip().rstrip("  "))
        index += 1
    box = document.add_table(rows=len(metadata), cols=1)
    box.alignment = WD_TABLE_ALIGNMENT.CENTER
    box.autofit = False
    box.columns[0].width = Cm(15.8)
    for row, text in zip(box.rows, metadata):
        cell = row.cells[0]
        set_cell_shading(cell, LIGHT)
        set_cell_margins(cell, 38, 100, 38, 100)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        add_inline(paragraph, text)
        for run in paragraph.runs:
            run.font.name = "Arial"
            run.font.size = Pt(9.2)
    separator = document.add_paragraph()
    separator.paragraph_format.space_before = Pt(5)
    separator.paragraph_format.space_after = Pt(1)
    p_pr = separator._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "14")
    bottom.set(qn("w:color"), NAVY)
    borders.append(bottom)
    p_pr.append(borders)
    return index + 1


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows = []
    index = start
    while index < len(lines) and lines[index].lstrip().startswith("|"):
        cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            rows.append(cells)
        index += 1
    return rows, index


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    table.style = "Table Grid"
    for row_index, values in enumerate(rows):
        row = table.rows[row_index]
        prevent_row_split(row)
        if row_index == 0:
            repeat_table_header(row)
        for column_index, value in enumerate(values):
            cell = row.cells[column_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_shading(cell, NAVY if row_index == 0 else (LIGHT if row_index % 2 == 0 else "FFFFFF"))
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if column_index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            add_inline(paragraph, value)
            for run in paragraph.runs:
                run.font.name = "Arial"
                run.font.size = Pt(8.2)
                if row_index == 0:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_image(document: Document, markdown_path: Path, line: str) -> None:
    match = re.match(r"!\[[^]]*\]\(([^)]+)\)", line.strip())
    if not match:
        return
    image_path = (markdown_path.parent / match.group(1)).resolve()
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(1)
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.paragraph_format.keep_together = True
    paragraph.add_run().add_picture(str(image_path), width=Cm(15.7))


def convert(markdown_path: Path, output_path: Path) -> None:
    lines = markdown_path.read_text(encoding="utf-8").splitlines()
    document = Document()
    style_document(document)
    document.core_properties.title = "Clasificación visual de residuos electrónicos con ETID"
    document.core_properties.subject = "Informe técnico final del proyecto integrador"
    document.core_properties.author = "Sebastian Mejia Ortiz; Laura Mariana Jimenez Mienez; Luis Gonzalo Duran"
    document.core_properties.keywords = "ETID, residuos electrónicos, visión por computador, EfficientNet, HOG"
    index = add_title_block(document, lines)
    document.add_page_break()
    in_references = False

    while index < len(lines):
        stripped = lines[index].strip()
        if not stripped or stripped == "---":
            index += 1
            continue
        if stripped.startswith("|" ):
            table_rows, index = parse_table(lines, index)
            add_table(document, table_rows)
            continue
        if stripped.startswith("!["):
            add_image(document, markdown_path, stripped)
            index += 1
            continue
        if stripped.startswith("## "):
            heading_text = stripped[3:]
            in_references = heading_text == "Referencias"
            paragraph = document.add_paragraph(style="Heading 1")
            add_inline(paragraph, heading_text)
            index += 1
            continue
        if stripped.startswith("### "):
            paragraph = document.add_paragraph(style="Heading 2")
            add_inline(paragraph, stripped[4:])
            index += 1
            continue

        paragraph = document.add_paragraph()
        add_inline(paragraph, stripped.rstrip("  "))
        if in_references:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.left_indent = Cm(1.0)
            paragraph.paragraph_format.first_line_indent = Cm(-1.0)
            paragraph.paragraph_format.space_after = Pt(4)
            for run in paragraph.runs:
                run.font.size = Pt(9)
        elif stripped.startswith("**Figura"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.space_before = Pt(3)
            paragraph.paragraph_format.space_after = Pt(0)
        elif stripped.startswith("*Nota.*"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_after = Pt(3)
            for run in paragraph.runs:
                run.font.size = Pt(8.3)
        elif stripped.startswith("*") and stripped.endswith("*"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.space_after = Pt(1)
        index += 1

    # Documento en español para revisión ortográfica.
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            run._element.get_or_add_rPr().append(OxmlElement("w:lang"))
            run._element.rPr[-1].set(qn("w:val"), "es-CO")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("markdown", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    convert(args.markdown.resolve(), args.output.resolve())
    print(args.output.resolve())


if __name__ == "__main__":
    main()
